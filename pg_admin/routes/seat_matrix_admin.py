"""Admin: upload + browse the NEET-PG counselling Seat Matrix (founder 2026-09-30).

Upload the authority's announced seat matrix as an .xlsx (and, optionally, the
original PDF for users to view). Rows are stored VERBATIM in pg_seat_matrix —
college & course names are NOT matched to our master DB, so users get the official
data fast. Re-uploading the same (counselling body + academic year) is a clean
replace. True-admin gated.
"""
import io
import logging
import openpyxl
from flask import render_template, request, redirect, url_for, flash, Response
from db import get_db
from core.users import get_user
from core.auth import login_required

DEFAULT_BODY = 'All India MCC'
DEFAULT_YEAR = '2026-27'


def _require_admin():
    u = get_user()
    return u if (u and u.get('is_admin')) else None


def _s(v):
    return (str(v).strip() if v is not None else '')


def _match_headers(headers):
    """Map each column index → our field, by fuzzy header keyword. Tolerates the
    founder's Excel headers AND the raw PDF headers (e.g. 'Category as per MCC',
    'Corrected Course Name', 'Name of Medical College / Medical Institution')."""
    fields = {}
    for i, h in enumerate(headers):
        n = _s(h).lower()
        if not n:
            continue
        if 'seat' in n:
            fields[i] = 'seats'
        elif 'course' in n:
            fields[i] = 'course_name'
        elif 'code' in n:
            # must precede the 'college' check — "College Code" contains "college"
            fields[i] = 'college_code'
        elif 'category' in n or 'mcc' in n or 'all india' in n or 'govt' in n:
            fields[i] = 'category'
        elif 'college' in n or 'institut' in n or 'name of' in n:
            fields[i] = 'college_name'
        elif 'state' in n:
            fields[i] = 'state'
        elif n.startswith('sl') or n.startswith('sr') or n in ('no', 's.no', 's no', '#'):
            fields[i] = 'sl_no'
    return fields


def _parse_int(v):
    try:
        return int(float(str(v).replace(',', '').strip()))
    except Exception:
        return 0


def _read_seat_matrix(file_bytes):
    """Return (rows, skipped). Each row = dict of our columns. Raises ValueError
    with a human message when the file is unusable."""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception:
        raise ValueError("Could not open the file — please upload a valid .xlsx workbook.")
    ws = wb[wb.sheetnames[0]]
    headers = None
    fmap = {}
    rows = []
    skipped = 0
    for r in ws.iter_rows(values_only=True):
        if r is None:
            continue
        if headers is None:
            if any(c not in (None, '') for c in r):
                headers = list(r)
                fmap = _match_headers(headers)
                need = {'college_name', 'course_name', 'seats'}
                if not need.issubset(set(fmap.values())):
                    wb.close()
                    raise ValueError(
                        "Could not find the expected columns. The sheet needs at least "
                        "a college name, a course name and a seats column "
                        f"(found: {sorted(set(fmap.values()))}).")
            continue
        if not any(c not in (None, '') for c in r):
            continue
        rec = {'sl_no': None, 'college_code': '', 'state': '', 'college_name': '',
               'category': '', 'course_name': '', 'seats': 0}
        for i, field in fmap.items():
            if i >= len(r):
                continue
            val = r[i]
            if field == 'seats':
                rec['seats'] = _parse_int(val)
            elif field == 'sl_no':
                rec['sl_no'] = _parse_int(val) or None
            else:
                rec[field] = _s(val)
        # A usable row must at least name a college and a course.
        if not rec['college_name'] or not rec['course_name']:
            skipped += 1
            continue
        rows.append(rec)
    wb.close()
    if not rows:
        raise ValueError("No data rows found in the sheet.")
    return rows, skipped


@login_required
def seat_matrix_admin():
    if not _require_admin():
        flash('Access denied', 'error'); return redirect(url_for('dashboard'))
    # Browse filters (server-rendered, no JS needed).
    b_body = _s(request.args.get('b_body'))
    b_year = _s(request.args.get('b_year'))
    f_state = _s(request.args.get('state'))
    f_cat = _s(request.args.get('category'))
    f_course = _s(request.args.get('course'))
    f_q = _s(request.args.get('q'))
    try:
        page = max(1, int(request.args.get('page', 1)))
    except Exception:
        page = 1
    page_size = 50

    conn = get_db()
    sources = []
    browse = None
    try:
        sources = conn.execute(
            "SELECT counselling_body, academic_year, source_file_name, pdf_name, "
            "row_count, college_count, total_seats, uploaded_by, uploaded_at "
            "FROM pg_seat_matrix_source ORDER BY academic_year DESC, counselling_body ASC"
        ).fetchall()
        # Default the browse target to the first loaded matrix.
        if not (b_body and b_year) and sources:
            b_body = sources[0]['counselling_body']
            b_year = sources[0]['academic_year']
        if b_body and b_year:
            base = "FROM pg_seat_matrix WHERE counselling_body=? AND academic_year=?"
            bp = [b_body, b_year]
            states = [r['state'] for r in conn.execute(
                f"SELECT DISTINCT state {base} AND state<>'' ORDER BY state", bp).fetchall()]
            cats = [r['category'] for r in conn.execute(
                f"SELECT DISTINCT category {base} AND category<>'' ORDER BY category", bp).fetchall()]
            courses = [r['course_name'] for r in conn.execute(
                f"SELECT DISTINCT course_name {base} AND course_name<>'' ORDER BY course_name", bp).fetchall()]
            conds = ["counselling_body=?", "academic_year=?"]; params = [b_body, b_year]
            if f_state:
                conds.append("state=?"); params.append(f_state)
            if f_cat:
                conds.append("category=?"); params.append(f_cat)
            if f_course:
                conds.append("course_name=?"); params.append(f_course)
            if f_q:
                from pg_admin.routes.api import smart_name_clause
                fn, pn = smart_name_clause("college_name", f_q)
                conds.append(f"({fn} OR college_code ILIKE ?)"); params.extend(pn + [f"%{f_q}%"])
            where = " AND ".join(conds)
            total = conn.execute(f"SELECT COUNT(*) AS n FROM pg_seat_matrix WHERE {where}",
                                 params).fetchone()['n']
            seats = conn.execute(f"SELECT COALESCE(SUM(seats),0) AS s FROM pg_seat_matrix WHERE {where}",
                                 params).fetchone()['s']
            offset = (page - 1) * page_size
            rows = conn.execute(
                f"SELECT sl_no, college_code, state, college_name, category, course_name, seats "
                f"FROM pg_seat_matrix WHERE {where} "
                f"ORDER BY state ASC, college_name ASC, course_name ASC "
                f"LIMIT {page_size} OFFSET {offset}", params).fetchall()
            browse = {
                'body': b_body, 'year': b_year,
                'states': states, 'categories': cats, 'courses': courses,
                'rows': rows, 'total': total, 'seats': int(seats or 0),
                'page': page, 'page_size': page_size,
                'total_pages': (total + page_size - 1) // page_size if page_size else 1,
                'state': f_state, 'category': f_cat, 'course': f_course, 'q': f_q,
            }
    except Exception as e:
        logging.error("seat_matrix_admin: %s", e)
        try: conn.rollback()
        except Exception: pass
    finally:
        conn.close()
    from pg_admin.authorities import all_authorities
    return render_template('pg_admin/seat_matrix.html', sources=sources, browse=browse,
                           default_body=DEFAULT_BODY, default_year=DEFAULT_YEAR,
                           authorities=all_authorities(),
                           active_section='goocampus_in')


@login_required
def seat_matrix_upload():
    if not _require_admin():
        flash('Access denied', 'error'); return redirect(url_for('dashboard'))
    body = _s(request.form.get('counselling_body')) or DEFAULT_BODY
    year = _s(request.form.get('academic_year')) or DEFAULT_YEAR
    # Explicit authority link (founder 2026-10-06). When chosen, it drives the body label so
    # the matrix ties to the authority by its stable code, not spelling.
    authority_code = _s(request.form.get('authority_code'))
    if authority_code:
        try:
            from pg_admin.authorities import get_authority
            _a = get_authority(authority_code)
            if _a:
                authority_code = _a['code']
                body = 'All India MCC' if _a['code'] == 'mcc' else _a['name']
        except Exception:
            pass
    data_file = request.files.get('data_file')
    pdf_file = request.files.get('pdf_file')

    if not data_file or not data_file.filename:
        flash('Please choose the seat-matrix Excel (.xlsx) file.', 'error')
        return redirect(url_for('pg_seat_matrix_admin'))
    if not data_file.filename.lower().endswith(('.xlsx', '.xlsm')):
        flash('The data file must be an Excel .xlsx workbook.', 'error')
        return redirect(url_for('pg_seat_matrix_admin'))

    try:
        rows, skipped = _read_seat_matrix(data_file.read())
    except ValueError as e:
        flash(str(e), 'error')
        return redirect(url_for('pg_seat_matrix_admin'))
    except Exception as e:
        logging.error("seat_matrix parse: %s", e)
        flash('Could not read the file. Please check it is the announced seat matrix.', 'error')
        return redirect(url_for('pg_seat_matrix_admin'))

    total_seats = sum(r['seats'] for r in rows)
    college_count = len({(r['college_code'], r['college_name']) for r in rows})

    # Optional PDF (kept if omitted on a re-upload).
    pdf_bytes = None; pdf_name = ''; pdf_ctype = 'application/pdf'
    if pdf_file and pdf_file.filename:
        if not pdf_file.filename.lower().endswith('.pdf'):
            flash('The official document must be a .pdf file.', 'error')
            return redirect(url_for('pg_seat_matrix_admin'))
        pdf_bytes = pdf_file.read()
        pdf_name = pdf_file.filename
        pdf_ctype = pdf_file.mimetype or 'application/pdf'

    user = get_user() or {}
    uploaded_by = user.get('name') or user.get('emp_code') or 'admin'

    conn = get_db()
    try:
        # Clean replace for this body + year.
        conn.execute("DELETE FROM pg_seat_matrix WHERE counselling_body=? AND academic_year=?",
                     (body, year))
        cols = ['counselling_body', 'academic_year', 'authority_code', 'sl_no', 'college_code',
                'state', 'college_name', 'category', 'course_name', 'seats']
        ph = ','.join(['?'] * len(cols))
        sql = f"INSERT INTO pg_seat_matrix ({','.join(cols)}) VALUES ({ph})"
        batch = [[body, year, authority_code, r['sl_no'], r['college_code'], r['state'],
                  r['college_name'], r['category'], r['course_name'], r['seats']] for r in rows]
        conn.execute_batch(sql, batch, page_size=1000)

        # Upsert the source row (keep an existing PDF if none supplied now).
        existing = conn.execute(
            "SELECT id FROM pg_seat_matrix_source WHERE counselling_body=? AND academic_year=?",
            (body, year)).fetchone()
        if existing:
            if pdf_bytes is not None:
                conn.execute(
                    "UPDATE pg_seat_matrix_source SET authority_code=?, source_file_name=?, pdf_name=?, "
                    "pdf_data=?, pdf_content_type=?, row_count=?, college_count=?, total_seats=?, "
                    "uploaded_by=?, uploaded_at=CURRENT_TIMESTAMP WHERE id=?",
                    (authority_code, data_file.filename, pdf_name, pdf_bytes, pdf_ctype,
                     len(rows), college_count, total_seats, uploaded_by, existing['id']))
            else:
                conn.execute(
                    "UPDATE pg_seat_matrix_source SET authority_code=?, source_file_name=?, row_count=?, "
                    "college_count=?, total_seats=?, uploaded_by=?, uploaded_at=CURRENT_TIMESTAMP "
                    "WHERE id=?",
                    (authority_code, data_file.filename, len(rows), college_count, total_seats,
                     uploaded_by, existing['id']))
        else:
            conn.execute(
                "INSERT INTO pg_seat_matrix_source (counselling_body, academic_year, authority_code, "
                "source_file_name, pdf_name, pdf_data, pdf_content_type, row_count, college_count, "
                "total_seats, uploaded_by) VALUES (?,?,?,?,?,?,?,?,?,?,?)",
                (body, year, authority_code, data_file.filename, pdf_name, pdf_bytes, pdf_ctype,
                 len(rows), college_count, total_seats, uploaded_by))
        conn.commit()
        msg = (f"Loaded {len(rows):,} rows · {college_count:,} colleges · {total_seats:,} seats "
               f"for {body} {year}.")
        if skipped:
            msg += f" ({skipped} blank/header rows skipped.)"
        if pdf_bytes is not None:
            msg += " Official PDF stored."
        flash(msg, 'success')
    except Exception as e:
        try: conn.rollback()
        except Exception: pass
        logging.error("seat_matrix_upload: %s", e)
        flash('Upload failed while saving. Please try again.', 'error')
    finally:
        conn.close()
    return redirect(url_for('pg_seat_matrix_admin'))


@login_required
def seat_matrix_delete():
    if not _require_admin():
        flash('Access denied', 'error'); return redirect(url_for('dashboard'))
    body = _s(request.form.get('counselling_body'))
    year = _s(request.form.get('academic_year'))
    if not body or not year:
        flash('Nothing to delete.', 'error')
        return redirect(url_for('pg_seat_matrix_admin'))
    conn = get_db()
    try:
        conn.execute("DELETE FROM pg_seat_matrix WHERE counselling_body=? AND academic_year=?",
                     (body, year))
        conn.execute("DELETE FROM pg_seat_matrix_source WHERE counselling_body=? AND academic_year=?",
                     (body, year))
        conn.commit()
        flash(f'Removed the {body} {year} seat matrix.', 'success')
    except Exception as e:
        try: conn.rollback()
        except Exception: pass
        logging.error("seat_matrix_delete: %s", e)
        flash('Could not delete.', 'error')
    finally:
        conn.close()
    return redirect(url_for('pg_seat_matrix_admin'))


@login_required
def seat_matrix_pdf_admin():
    """Admin preview of the stored official PDF."""
    if not _require_admin():
        flash('Access denied', 'error'); return redirect(url_for('dashboard'))
    body = _s(request.args.get('body')) or DEFAULT_BODY
    year = _s(request.args.get('year')) or DEFAULT_YEAR
    conn = get_db()
    try:
        row = conn.execute(
            "SELECT pdf_name, pdf_data, pdf_content_type FROM pg_seat_matrix_source "
            "WHERE counselling_body=? AND academic_year=?", (body, year)).fetchone()
    finally:
        conn.close()
    if not row or not row['pdf_data']:
        flash('No PDF stored for that seat matrix.', 'error')
        return redirect(url_for('pg_seat_matrix_admin'))
    fname = row['pdf_name'] or 'seat-matrix.pdf'
    return Response(bytes(row['pdf_data']), mimetype=row['pdf_content_type'] or 'application/pdf',
                    headers={'Content-Disposition': f'inline; filename="{fname}"'})
