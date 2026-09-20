"""Admin: import + browse the unified college master (founder 2026-09-18).

Upload the College DB workbook (Medical & Healthcare + DNB Hospitals sheets) →
pg_college_master (one row per college, per kind) + pg_college_course (course
catalogue) + a canonical alias per college. Then upload the audited matching
list → cut-off-name aliases so the already-uploaded pg_cutoffs (incl. stipend/
bond/penalty) resolve to one master college. True-admin gated. Re-uploading is
a clean rebuild (workbook first, then matching).
"""
import io
import re
import logging
import openpyxl
from flask import render_template, request, redirect, url_for, flash
from db import get_db
from core.users import get_user
from core.auth import login_required

_MASTER_COLS = {
    'College Name': 'college_name', 'University': 'university', 'Logo URL': 'logo_url',
    'City': 'city', 'State': 'state', 'District': 'district', 'Country': 'country',
    'College Type': 'college_type', 'Accreditation': 'accreditation', 'College Stream': 'college_stream',
    'Established Year': 'established_year', 'Nearest Airport': 'nearest_airport',
    'Winter Min Temp': 'winter_min_temp', 'Summer Max Temp': 'summer_max_temp',
    'Latitude': 'latitude', 'Longitude': 'longitude',
    'Mess Fee Currency': 'mess_fee_currency', 'Hostel Fee Currency': 'hostel_fee_currency',
    'OPD': 'opd', 'IPD': 'ipd', 'Bed Count': 'bed_count', 'Official Website': 'official_website',
}
_MASTER_NUM = {'Mess Fee (Min)': 'mess_fee_min', 'Mess Fee (Max)': 'mess_fee_max',
               'Hostel Fee (Min)': 'hostel_fee_min', 'Hostel Fee (Max)': 'hostel_fee_max'}
_COURSE_COLS = {'Course': 'course', 'Course Level': 'course_level', 'Course Stream': 'course_stream',
    'Seat Intake': 'seat_intake', 'Exam Type': 'exam_type', 'Entrance Exams': 'entrance_exams',
    'Entrance Exam Eligibility': 'entrance_exam_eligibility', 'Academic Eligibility': 'academic_eligibility',
    'Duration (Total Months)': 'duration_total_months', 'Duration (Years)': 'duration_years',
    'Duration (Months)': 'duration_months'}
_SHEET_KIND = {'Medical & Healthcare': 'medical', 'DNB Hospitals': 'dnb'}

# Cut-off → DB-name links NOT in the matching file (founder-confirmed 2026-09-18).
_SPECIAL_CUTOFF = {
    'Government Institute of Medical Sciences, Greater Noida':
        'Government Institute of Medical Sciences, Kasna, Greater Noida',
    'Autonomous State Medical College, Ayodhya':
        'Government Medical College, Faizabad (Ayodhya)',
}

# Fingerprint of the cut-off master file the founder shared, for the audit compare
# ("GooCampus NEET PG 2025 - MASTER only.xlsx"; computed 2026-09-18).
_CUTOFF_FILE = {
    'name': 'GooCampus NEET PG 2025 - MASTER only.xlsx',
    'rows': 37666, 'institutes': 847, 'states': 33, 'courses': 99, 'year': '2025',
    'stipend': 37609, 'bond': 37189, 'penalty': 37609,
    'med_institutes': 721, 'dnb_institutes': 154,   # by degree family (28 in both)
}


def _s(v):
    return (str(v).strip() if v is not None else '')


def _norm(v):
    return re.sub(r'[^a-z0-9]+', ' ', _s(v).lower()).strip()


def _num(v):
    try:
        return float(str(v).replace(',', '').strip()) if _s(v) else None
    except Exception:
        return None


# Course buckets for the College Database tabs (founder 2026-09-18).
_CAT_LABELS = [('mbbs', 'MBBS (UG)'), ('mdms', 'MD / MS'),
               ('super', 'Super Speciality'), ('diploma', 'Diploma'), ('dnb', 'DNB')]
_CAT_KEYS = [k for k, _l in _CAT_LABELS]


def _course_cat(course, level, kind):
    """Bucket a course row: mbbs | mdms | super | diploma | dnb."""
    c = _s(course).upper()
    lv = _s(level).lower()
    if kind == 'dnb':
        return 'dnb'
    if 'under grad' in lv or c.startswith('MBBS'):
        return 'mbbs'
    if c.startswith('M.CH') or c.startswith('MCH') or c.startswith('DM ') or c.startswith('DM-') or c.startswith('DM.'):
        return 'super'
    if 'diploma' in lv or c.startswith('DIPLOMA'):
        return 'diploma'
    return 'mdms'   # MD / MS and other PG-medical


def _require_admin():
    u = get_user()
    return u if (u and u.get('is_admin')) else None


def _stats(conn):
    st = {}
    def c(q, *p):
        try: return conn.execute(q, p).fetchone()['n']
        except Exception: return 0
    st['medical'] = c("SELECT COUNT(*) AS n FROM pg_college_master WHERE kind='medical'")
    st['dnb'] = c("SELECT COUNT(*) AS n FROM pg_college_master WHERE kind='dnb'")
    st['master'] = st['medical'] + st['dnb']
    st['courses'] = c("SELECT COUNT(*) AS n FROM pg_college_course")
    st['aliases'] = c("SELECT COUNT(*) AS n FROM pg_college_alias")
    st['cutoff_aliases'] = c("SELECT COUNT(*) AS n FROM pg_college_alias WHERE alias_source='cutoff'")
    st['cutoff_colleges'] = c("SELECT COUNT(DISTINCT institute) AS n FROM pg_cutoffs "
                              "WHERE COALESCE(institute,'') <> ''")
    return st




def _ensure_log(conn):
    conn.execute('''CREATE TABLE IF NOT EXISTS pg_college_import_log (
        id SERIAL PRIMARY KEY, job TEXT, status TEXT DEFAULT 'running',
        detail TEXT DEFAULT '', started_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
        finished_at TIMESTAMP)''')


def _log_start(job):
    """Insert a 'running' log row, return its id (own connection)."""
    conn = get_db()
    try:
        _ensure_log(conn)
        lid = conn.execute("INSERT INTO pg_college_import_log (job, status) VALUES (?, 'running') RETURNING id",
                           [job]).fetchone()['id']
        conn.commit()
        return lid
    finally:
        try: conn.close()
        except Exception: pass


def _log_finish(conn, lid, status, detail):
    conn.execute("UPDATE pg_college_import_log SET status=?, detail=?, finished_at=CURRENT_TIMESTAMP WHERE id=?",
                 [status, detail[:1000], lid])
    conn.commit()


def _latest_log(conn, job):
    try:
        return conn.execute("SELECT * FROM pg_college_import_log WHERE job=? ORDER BY id DESC LIMIT 1",
                            [job]).fetchone()
    except Exception:
        return None


@login_required
def college_master_admin():
    u = _require_admin()
    if not u:
        flash('Access denied', 'error'); return redirect(url_for('dashboard'))
    conn = get_db()
    try:
        _ensure_log(conn); conn.commit()
        stats = _stats(conn)
        wb_log = _latest_log(conn, 'workbook')
        mt_log = _latest_log(conn, 'matching')
    finally:
        conn.close()
    running = (wb_log and wb_log['status'] == 'running') or (mt_log and mt_log['status'] == 'running')
    return render_template('pg_admin/college_master.html', stats=stats,
                           wb_log=wb_log, mt_log=mt_log, running=running,
                           active_section='goocampus_in')


# ── Workbook import (runs in a background thread) ─────────────────────────────
def _run_workbook(data, lid):
    conn = get_db()
    try:
        master_attrs = list(_MASTER_COLS.values()) + list(_MASTER_NUM.values())
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        masters = {}       # (kind, master_key) -> [name, *fixed-order values]
        courses_raw = []   # ((kind, master_key), [course values])
        for sheet, kind in _SHEET_KIND.items():
            if sheet not in wb.sheetnames:
                continue
            rows = wb[sheet].iter_rows(values_only=True)
            idx = {h: i for i, h in enumerate([_s(h) for h in next(rows)])}
            for r in rows:
                if not r:
                    continue
                name = _s(r[idx['College Name']]) if 'College Name' in idx else ''
                if not name:
                    continue
                state = _s(r[idx['State']]) if 'State' in idx else ''
                gkey = (kind, _norm(name) + '|' + _norm(state))
                if gkey not in masters:
                    vals = [name]
                    for hdr in _MASTER_COLS:
                        vals.append(_s(r[idx[hdr]]) if hdr in idx else '')
                    for hdr in _MASTER_NUM:
                        vals.append(_num(r[idx[hdr]]) if hdr in idx else None)
                    masters[gkey] = vals
                cv = [_s(r[idx[hdr]]) if hdr in idx else '' for hdr in _COURSE_COLS]
                if any(cv):
                    cv.append(_course_cat(cv[0], cv[1], kind))   # cv[0]=course, cv[1]=course_level
                    courses_raw.append((gkey, cv))
        wb.close()

        conn.execute("DELETE FROM pg_college_alias")
        conn.execute("DELETE FROM pg_college_course")
        conn.execute("DELETE FROM pg_college_master")

        mcols = 'kind,master_key,' + ','.join(master_attrs)
        mph = ','.join(['?'] * (2 + len(master_attrs)))
        master_batch = [[k[0], k[1]] + v[1:] for k, v in masters.items()]
        conn.execute_batch(f"INSERT INTO pg_college_master ({mcols}) VALUES ({mph})",
                           master_batch, page_size=1000)

        idmap = {(row['kind'], row['master_key']): row['id']
                 for row in conn.execute("SELECT id, kind, master_key FROM pg_college_master").fetchall()}

        alias_rows = [(idmap[k], v[0], _norm(v[0]), 'canonical') for k, v in masters.items()]
        conn.execute_batch(
            "INSERT INTO pg_college_alias (master_id, alias_name, alias_key, alias_source) VALUES (?,?,?,?)",
            alias_rows, page_size=1000)

        course_batch = [[idmap[gkey]] + cv for gkey, cv in courses_raw]
        ccols = 'master_id,' + ','.join(_COURSE_COLS.values()) + ',course_category'
        cph = ','.join(['?'] * (2 + len(_COURSE_COLS)))
        conn.execute_batch(f"INSERT INTO pg_college_course ({ccols}) VALUES ({cph})",
                           course_batch, page_size=1000)
        conn.commit()
        _log_finish(conn, lid, 'done',
                    f'Imported {len(master_batch)} colleges + {len(course_batch)} courses. '
                    'Now upload the matching list to link cut-offs.')
    except Exception as e:
        try: conn.rollback()
        except Exception: pass
        logging.error(f"_run_workbook: {e}")
        try: _log_finish(conn, lid, 'error', f'Import failed: {e}')
        except Exception: pass
    finally:
        try: conn.close()
        except Exception: pass


@login_required
def college_master_upload_workbook():
    u = _require_admin()
    if not u:
        flash('Access denied', 'error'); return redirect(url_for('dashboard'))
    fs = request.files.get('workbook')
    if not fs or not fs.filename:
        flash('Choose the college DB workbook (.xlsx) first.', 'error')
        return redirect(url_for('pg_college_master_admin'))
    data = fs.read()
    lid = _log_start('workbook')
    import threading
    threading.Thread(target=_run_workbook, args=(data, lid), daemon=True).start()
    flash('Import started — reading the workbook and loading colleges. This page refreshes '
          'itself; the result will show here in a few seconds.', 'info')
    return redirect(url_for('pg_college_master_admin'))


# ── Matching / cut-off linker (runs in a background thread) ───────────────────
def _run_matching(data, lid):
    conn = get_db()
    try:
        name2mid = {}   # canonical norm(name) -> master_id (prefer medical)
        for row in conn.execute(
                "SELECT a.master_id, a.alias_key, m.kind FROM pg_college_alias a "
                "JOIN pg_college_master m ON m.id = a.master_id "
                "WHERE a.alias_source='canonical'").fetchall():
            k = row['alias_key']
            if k not in name2mid or row['kind'] == 'medical':
                name2mid[k] = row['master_id']
        if not name2mid:
            _log_finish(conn, lid, 'error', 'Upload the college DB workbook first.')
            return

        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        rows = list(wb.worksheets[0].iter_rows(values_only=True)); wb.close()
        hr = next(i for i, r in enumerate(rows[:6]) if r and any(_s(c) == 'Match Type' for c in r))
        hdr = [_s(c) for c in rows[hr]]
        ci = hdr.index('Cut-off (Master) College Name'); di = hdr.index('DB Matched College Name')
        pairs = {}
        for r in rows[hr + 1:]:
            if r and _s(r[ci]):
                pairs[_s(r[ci])] = _s(r[di])
        for cut, db in _SPECIAL_CUTOFF.items():
            pairs.setdefault(cut, db)

        cut_names = [r['institute'] for r in
                     conn.execute("SELECT DISTINCT institute FROM pg_cutoffs "
                                  "WHERE COALESCE(institute,'') <> ''").fetchall()]
        conn.execute("DELETE FROM pg_college_alias WHERE alias_source='cutoff'")
        linked, leftover = [], []
        for cut in cut_names:
            db = pairs.get(cut, cut)
            mid = name2mid.get(_norm(db)) or name2mid.get(_norm(cut))
            if mid:
                linked.append((mid, cut, _norm(cut), 'cutoff'))
            else:
                leftover.append(cut)
        if linked:
            conn.execute_batch(
                "INSERT INTO pg_college_alias (master_id, alias_name, alias_key, alias_source) VALUES (?,?,?,?)",
                linked, page_size=1000)
        conn.commit()
        msg = f'Linked {len(linked)} of {len(cut_names)} cut-off colleges to the master.'
        if leftover:
            msg += ' Unlinked: ' + '; '.join(leftover[:8]) + (' …' if len(leftover) > 8 else '')
        else:
            msg += ' 100% linked ✓'
        _log_finish(conn, lid, 'done', msg)
    except Exception as e:
        try: conn.rollback()
        except Exception: pass
        logging.error(f"_run_matching: {e}")
        try: _log_finish(conn, lid, 'error', f'Matching import failed: {e}')
        except Exception: pass
    finally:
        try: conn.close()
        except Exception: pass


@login_required
def college_master_upload_matching():
    u = _require_admin()
    if not u:
        flash('Access denied', 'error'); return redirect(url_for('dashboard'))
    fs = request.files.get('matching')
    if not fs or not fs.filename:
        flash('Choose the matching list (.xlsx) first.', 'error')
        return redirect(url_for('pg_college_master_admin'))
    data = fs.read()
    lid = _log_start('matching')
    import threading
    threading.Thread(target=_run_matching, args=(data, lid), daemon=True).start()
    flash('Linking started — matching cut-off colleges to the master. This page refreshes '
          'itself; the result will show here in a few seconds.', 'info')
    return redirect(url_for('pg_college_master_admin'))


# ── Read-only audit: compare the loaded cut-offs vs the shared master file ─────
@login_required
def college_master_cutoff_audit():
    u = _require_admin()
    if not u:
        flash('Access denied', 'error'); return redirect(url_for('dashboard'))
    conn = get_db()
    db = {}
    years = []
    try:
        def one(q):
            try:
                return conn.execute(q).fetchone()['n']
            except Exception:
                try: conn.rollback()
                except Exception: pass
                return None
        db['rows'] = one("SELECT COUNT(*) AS n FROM pg_cutoffs")
        # Rows that carry an actual cut-off rank in ANY round (R1-R4 / Stray / closing).
        _rank = ("(r1 IS NOT NULL OR r2 IS NOT NULL OR r3 IS NOT NULL OR r4 IS NOT NULL "
                 "OR stray IS NOT NULL OR closing_rank IS NOT NULL)")
        db['with_rank'] = one(f"SELECT COUNT(*) AS n FROM pg_cutoffs WHERE {_rank}")
        db['no_rank'] = one(f"SELECT COUNT(*) AS n FROM pg_cutoffs WHERE NOT {_rank}")
        db['is_reference'] = one("SELECT COUNT(*) AS n FROM pg_cutoffs WHERE COALESCE(is_reference,0)=1")
        db['institutes'] = one("SELECT COUNT(DISTINCT institute) AS n FROM pg_cutoffs WHERE COALESCE(institute,'') <> ''")
        # Stipend section: distinct institutes by degree family (mirrors the two tabs).
        db['med_institutes'] = one("SELECT COUNT(DISTINCT institute) AS n FROM pg_cutoffs "
                                   "WHERE UPPER(COALESCE(degree,'')) NOT LIKE '%%DNB%%' AND COALESCE(institute,'') <> ''")
        db['dnb_institutes'] = one("SELECT COUNT(DISTINCT institute) AS n FROM pg_cutoffs "
                                   "WHERE UPPER(COALESCE(degree,'')) LIKE '%%DNB%%' AND COALESCE(institute,'') <> ''")
        # Of the blank-rank rows: how many carry stipend/bond/penalty, and how many
        # belong to a college that has NO ranked row anywhere (its data lives ONLY here).
        _money = ("(stipend IS NOT NULL OR stipend_yr2 IS NOT NULL OR stipend_yr3 IS NOT NULL "
                  "OR bond_years IS NOT NULL OR penalty IS NOT NULL)")
        db['blank_with_money'] = one(f"SELECT COUNT(*) AS n FROM pg_cutoffs WHERE NOT {_rank} AND {_money}")
        db['blank_only_colleges'] = one(
            f"SELECT COUNT(*) AS n FROM ("
            f"  SELECT DISTINCT institute FROM pg_cutoffs WHERE NOT {_rank} AND COALESCE(institute,'')<>'' "
            f"  EXCEPT "
            f"  SELECT DISTINCT institute FROM pg_cutoffs WHERE {_rank}) t")
        try:
            blank_only_names = [r['institute'] for r in conn.execute(
                f"SELECT DISTINCT institute FROM pg_cutoffs WHERE NOT {_rank} AND COALESCE(institute,'')<>'' "
                f"EXCEPT SELECT DISTINCT institute FROM pg_cutoffs WHERE {_rank} LIMIT 15").fetchall()]
        except Exception:
            conn.rollback(); blank_only_names = []
        db['states'] = one("SELECT COUNT(DISTINCT state) AS n FROM pg_cutoffs WHERE COALESCE(state,'') <> ''")
        db['courses'] = one("SELECT COUNT(DISTINCT course) AS n FROM pg_cutoffs WHERE COALESCE(course,'') <> ''")
        db['stipend'] = one("SELECT COUNT(*) AS n FROM pg_cutoffs WHERE stipend IS NOT NULL")
        db['bond'] = one("SELECT COUNT(*) AS n FROM pg_cutoffs WHERE bond_years IS NOT NULL")
        db['penalty'] = one("SELECT COUNT(*) AS n FROM pg_cutoffs WHERE penalty IS NOT NULL")
        try:
            years = [dict(r) for r in conn.execute(
                "SELECT year, COUNT(*) AS n FROM pg_cutoffs GROUP BY year ORDER BY year").fetchall()]
        except Exception:
            conn.rollback()
        # cut-offs whose institute name did NOT resolve to a master college
        db['linked'] = one("SELECT COUNT(DISTINCT c.institute) AS n FROM pg_cutoffs c "
                           "JOIN pg_college_alias a ON a.alias_key = "
                           "btrim(regexp_replace(lower(c.institute), '[^a-z0-9]+', ' ', 'g')) "
                           "WHERE COALESCE(c.institute,'') <> ''")
    finally:
        conn.close()

    rows_cmp = [
        ('Cut-off rows (with a rank in any round)', _CUTOFF_FILE['rows'], db.get('with_rank')),
        ('Distinct colleges',  _CUTOFF_FILE['institutes'], db.get('institutes')),
        ('Distinct states',    _CUTOFF_FILE['states'], db.get('states')),
        ('Distinct courses',   _CUTOFF_FILE['courses'], db.get('courses')),
        ('Stipend — MD/MS colleges', _CUTOFF_FILE['med_institutes'], db.get('med_institutes')),
        ('Stipend — DNB colleges',   _CUTOFF_FILE['dnb_institutes'], db.get('dnb_institutes')),
    ]
    # The real question: do the RANKED rows match the file? Extra blank-rank rows
    # (stipend/bond/penalty-only) are expected and explain any total-row gap.
    verdict = all(exp == got for _l, exp, got in rows_cmp)
    total_gap = (db.get('rows') or 0) - (db.get('with_rank') or 0)
    return render_template('pg_admin/college_cutoff_audit.html',
                           file=_CUTOFF_FILE, rows_cmp=rows_cmp, years=years,
                           db=db, verdict=verdict, total_gap=total_gap,
                           blank_only_names=blank_only_names,
                           active_section='goocampus_in')


# ── Purge the blank-rank cut-off rows (destructive; server-side safety re-check) ──
@login_required
def college_master_purge_blank():
    u = _require_admin()
    if not u:
        flash('Access denied', 'error'); return redirect(url_for('dashboard'))
    conn = get_db()
    try:
        _rank = ("(r1 IS NOT NULL OR r2 IS NOT NULL OR r3 IS NOT NULL OR r4 IS NOT NULL "
                 "OR stray IS NOT NULL OR closing_rank IS NOT NULL)")
        # SAFETY: refuse if any college would lose all its data (exists only as blank rows).
        blank_only = conn.execute(
            f"SELECT COUNT(*) AS n FROM ("
            f"  SELECT DISTINCT institute FROM pg_cutoffs WHERE NOT {_rank} AND COALESCE(institute,'')<>'' "
            f"  EXCEPT SELECT DISTINCT institute FROM pg_cutoffs WHERE {_rank}) t").fetchone()['n']
        if blank_only and blank_only > 0:
            conn.rollback()
            flash(f'Refused — {blank_only} college(s) exist only as blank-rank rows and would lose '
                  'their stipend/bond/penalty. Nothing was deleted.', 'error')
            return redirect(url_for('pg_college_master_cutoff_audit'))
        before = conn.execute("SELECT COUNT(*) AS n FROM pg_cutoffs").fetchone()['n']
        conn.execute(f"DELETE FROM pg_cutoffs WHERE NOT {_rank}")
        after = conn.execute("SELECT COUNT(*) AS n FROM pg_cutoffs").fetchone()['n']
        conn.commit()
        flash(f'Removed {before - after} blank-rank rows. Cut-off dataset is now {after:,} rows '
              '(all with a rank).', 'success')
    except Exception as e:
        try: conn.rollback()
        except Exception: pass
        logging.error(f"college_master_purge_blank: {e}")
        flash(f'Delete failed: {e}', 'error')
    finally:
        try: conn.close()
        except Exception: pass
    return redirect(url_for('pg_college_master_cutoff_audit'))


# ── College Database browser (view-only) ─────────────────────────────────────
def _ensure_course_categories(conn):
    """Backfill course_category for rows imported before the column existed.
    One set-based UPDATE (fast); falls back to a row loop only on error. The CASE
    mirrors _course_cat(); % are doubled because the db shim passes params."""
    try:
        pending = conn.execute("SELECT 1 FROM pg_college_course "
                               "WHERE COALESCE(course_category,'') = '' LIMIT 1").fetchone()
        if not pending:
            return
    except Exception:
        conn.rollback(); return
    try:
        conn.execute(
            "UPDATE pg_college_course AS c SET course_category = CASE "
            "  WHEN m.kind = 'dnb' THEN 'dnb' "
            "  WHEN c.course_level ILIKE '%%under grad%%' OR c.course ILIKE 'MBBS%%' THEN 'mbbs' "
            "  WHEN c.course ILIKE 'M.CH%%' OR c.course ILIKE 'MCH%%' OR c.course ILIKE 'DM %%' "
            "       OR c.course ILIKE 'DM-%%' OR c.course ILIKE 'DM.%%' THEN 'super' "
            "  WHEN c.course_level ILIKE '%%diploma%%' OR c.course ILIKE 'DIPLOMA%%' THEN 'diploma' "
            "  ELSE 'mdms' END "
            "FROM pg_college_master AS m "
            "WHERE m.id = c.master_id AND COALESCE(c.course_category,'') = ''")
        conn.commit()
    except Exception as e:
        conn.rollback()
        logging.error(f"_ensure_course_categories set-update failed, falling back: {e}")
        try:
            todo = conn.execute(
                "SELECT c.id, c.course, c.course_level, m.kind FROM pg_college_course c "
                "JOIN pg_college_master m ON m.id = c.master_id "
                "WHERE COALESCE(c.course_category,'') = ''").fetchall()
            updates = [(_course_cat(r['course'], r['course_level'], r['kind']), r['id']) for r in todo]
            conn.executemany("UPDATE pg_college_course SET course_category = ? WHERE id = ?", updates)
            conn.commit()
        except Exception:
            conn.rollback()


@login_required
def college_database_list():
    u = _require_admin()
    if not u:
        flash('Access denied', 'error'); return redirect(url_for('dashboard'))
    q = _s(request.args.get('q'))
    cat = _s(request.args.get('cat'))
    state = _s(request.args.get('state'))
    conn = get_db()
    rows, states, total, cat_counts = [], [], 0, {}
    try:
        _ensure_course_categories(conn)
        where, params = [], []
        if q:
            where.append("m.college_name ILIKE ?"); params.append('%' + q + '%')
        if cat in _CAT_KEYS:
            where.append("EXISTS (SELECT 1 FROM pg_college_course c "
                         "WHERE c.master_id = m.id AND c.course_category = ?)")
            params.append(cat)
        if state:
            where.append("m.state = ?"); params.append(state)
        wsql = (' WHERE ' + ' AND '.join(where)) if where else ''
        total = conn.execute(f"SELECT COUNT(*) AS n FROM pg_college_master m{wsql}", params).fetchone()['n']
        rows = conn.execute(
            f"SELECT m.id, m.college_name, m.kind, m.city, m.state, m.college_type, m.logo_url, "
            f"(SELECT COUNT(*) FROM pg_college_course c WHERE c.master_id = m.id) AS n_courses "
            f"FROM pg_college_master m{wsql} ORDER BY m.college_name LIMIT 400", params).fetchall()
        states = [r['state'] for r in conn.execute(
            "SELECT DISTINCT state FROM pg_college_master WHERE COALESCE(state,'') <> '' ORDER BY state").fetchall()]
        all_names = [r['college_name'] for r in conn.execute(
            "SELECT college_name FROM pg_college_master ORDER BY college_name").fetchall()]
        # colleges offering ≥1 course in each bucket
        for r in conn.execute("SELECT course_category AS cat, COUNT(DISTINCT master_id) AS n "
                              "FROM pg_college_course GROUP BY course_category").fetchall():
            cat_counts[r['cat']] = r['n']
        n_total = conn.execute("SELECT COUNT(*) AS n FROM pg_college_master").fetchone()['n']
        # which buckets each shown college offers (for card chips)
        order = {k: i for i, (k, _l) in enumerate(_CAT_LABELS)}
        cat_by_college = {}
        ids = [r['id'] for r in rows]
        if ids:
            ph = ','.join(['?'] * len(ids))
            for r in conn.execute(f"SELECT DISTINCT master_id, course_category FROM pg_college_course "
                                  f"WHERE master_id IN ({ph})", ids).fetchall():
                cat_by_college.setdefault(r['master_id'], set()).add(r['course_category'])
        cat_by_college = {k: sorted([c for c in v if c in order], key=lambda c: order[c])
                          for k, v in cat_by_college.items()}
    finally:
        conn.close()
    return render_template('pg_admin/college_database.html', rows=rows, total=total,
                           states=states, q=q, cat=cat, state=state, all_names=all_names,
                           cat_labels=_CAT_LABELS, cat_counts=cat_counts, n_total=n_total,
                           cat_by_college=cat_by_college, cat_label_map=dict(_CAT_LABELS),
                           active_section='goocampus_in')


@login_required
def college_profile(master_id):
    u = _require_admin()
    if not u:
        flash('Access denied', 'error'); return redirect(url_for('dashboard'))
    conn = get_db()
    try:
        m = conn.execute("SELECT * FROM pg_college_master WHERE id = ?", [master_id]).fetchone()
        if not m:
            flash('College not found', 'error')
            return redirect(url_for('pg_college_database_list'))
        courses = conn.execute(
            "SELECT * FROM pg_college_course WHERE master_id = ? ORDER BY course", [master_id]).fetchall()
        aliases = conn.execute(
            "SELECT alias_name, alias_source FROM pg_college_alias WHERE master_id = ? "
            "ORDER BY alias_source, alias_name", [master_id]).fetchall()
        names = [a['alias_name'] for a in aliases]
        cutoffs = []
        if names:
            ph = ','.join(['?'] * len(names))
            cutoffs = conn.execute(
                f"SELECT course, category, quota, seat_type, r1, r2, r3, r4, stray, closing_rank, "
                f"fee, stipend, stipend_yr2, stipend_yr3, bond_years, penalty "
                f"FROM pg_cutoffs WHERE institute IN ({ph}) "
                f"ORDER BY course, category, quota", names).fetchall()
    finally:
        conn.close()
    # Group cut-offs by course + pull a representative stipend/bond/penalty (college-level).
    cutoffs_by_course, money = {}, {}
    for c in cutoffs:
        cutoffs_by_course.setdefault(c['course'] or '—', []).append(c)
        for k in ('stipend', 'stipend_yr2', 'stipend_yr3', 'bond_years', 'penalty'):
            if k not in money and c[k] is not None:
                money[k] = c[k]
    # Non-canonical names only, for the "also known as" chips near the top.
    other_names = sorted({a['alias_name'] for a in aliases
                          if _norm(a['alias_name']) != _norm(m['college_name'])})
    return render_template('pg_admin/college_profile.html', m=m, courses=courses,
                           aliases=aliases, cutoffs=cutoffs,
                           cutoffs_by_course=cutoffs_by_course, money=money,
                           other_names=other_names, active_section='goocampus_in')


# ── Stipend · Bond · Penalty section ─────────────────────────────────────────
# Driven DIRECTLY by the cut-off data (pg_cutoffs), split by DEGREE family so it
# mirrors the Excel exactly: MD/MS side = degrees without 'DNB' (MD/MS/PG-Diploma/
# MCh/MPH); DNB side = degrees containing 'DNB' (DNB/DNB-Diploma). A college with
# both (e.g. a medical college with DNB-quota seats) shows in both tabs, each with
# only that family's courses.
def _stipend_family(kind):
    """Return (deg_sql_clause, deg_param) for the degree family filter."""
    if kind == 'dnb':
        return "UPPER(COALESCE(c.degree,'')) LIKE ?", '%DNB%'
    return "UPPER(COALESCE(c.degree,'')) NOT LIKE ?", '%DNB%'


_STIPEND_PER_PAGE = 100


@login_required
def college_stipend():
    u = _require_admin()
    if not u:
        flash('Access denied', 'error'); return redirect(url_for('dashboard'))
    kind = _s(request.args.get('kind'))
    if kind not in ('medical', 'dnb'):
        kind = 'medical'                    # always one of the two tabs (no "all")
    state = _s(request.args.get('state'))
    q = _s(request.args.get('q'))
    sort = _s(request.args.get('sort')) or 'name'
    order = {'name': 'college_name ASC',
             'stipend_desc': 's1_max DESC NULLS LAST, college_name ASC',
             'stipend_asc': 's1_max ASC NULLS LAST, college_name ASC'}.get(sort, 'college_name ASC')
    try:
        page = max(1, int(request.args.get('page', 1)))
    except Exception:
        page = 1
    deg_clause, deg_param = _stipend_family(kind)
    conn = get_db()
    rows, states, total, all_names = [], [], 0, []
    try:
        # Grouped by the MASTER college so the name is the SAME canonical name shown
        # everywhere (College Database + Stipend); driven by cut-off degree family.
        where = [deg_clause,
                 "(c.stipend IS NOT NULL OR c.bond_years IS NOT NULL OR c.penalty IS NOT NULL)"]
        params = [deg_param]
        if state:
            where.append("m.state = ?"); params.append(state)
        if q:
            where.append("m.college_name ILIKE ?"); params.append('%' + q + '%')
        frm = (" FROM pg_cutoffs c "
               "JOIN pg_college_alias a ON a.alias_key = btrim(regexp_replace(lower(c.institute), '[^a-z0-9]+', ' ', 'g')) "
               "JOIN pg_college_master m ON m.id = a.master_id "
               "WHERE " + " AND ".join(where) +
               " GROUP BY m.id, m.college_name, m.state, m.kind")
        total = conn.execute("SELECT COUNT(*) AS n FROM (SELECT m.id" + frm + ") t", params).fetchone()['n']
        offset = (page - 1) * _STIPEND_PER_PAGE
        rows = conn.execute(
            "SELECT m.id AS id, m.college_name AS name, m.state AS state, m.kind AS kind, "
            "MIN(c.stipend) AS s1_min, MAX(c.stipend) AS s1_max, "
            "MIN(c.stipend_yr2) AS s2_min, MAX(c.stipend_yr2) AS s2_max, "
            "MIN(c.stipend_yr3) AS s3_min, MAX(c.stipend_yr3) AS s3_max, "
            "MIN(c.bond_years) AS b_min, MAX(c.bond_years) AS b_max, "
            "MIN(c.penalty) AS p_min, MAX(c.penalty) AS p_max, "
            "COUNT(DISTINCT c.course) AS n_courses" + frm +
            " ORDER BY " + order + " LIMIT ? OFFSET ?", params + [_STIPEND_PER_PAGE, offset]).fetchall()
        states = [r['state'] for r in conn.execute(
            "SELECT DISTINCT state FROM pg_college_master WHERE COALESCE(state,'') <> '' ORDER BY state").fetchall()]
        # autocomplete names for THIS family (with stipend data), ignoring paging/filters
        all_names = [r['college_name'] for r in conn.execute(
            "SELECT DISTINCT m.college_name FROM pg_cutoffs c "
            "JOIN pg_college_alias a ON a.alias_key = btrim(regexp_replace(lower(c.institute), '[^a-z0-9]+', ' ', 'g')) "
            "JOIN pg_college_master m ON m.id = a.master_id "
            "WHERE " + deg_clause + " AND (c.stipend IS NOT NULL OR c.bond_years IS NOT NULL OR c.penalty IS NOT NULL) "
            "ORDER BY m.college_name", [deg_param]).fetchall()]
    finally:
        conn.close()
    pages = max(1, (total + _STIPEND_PER_PAGE - 1) // _STIPEND_PER_PAGE)
    return render_template('pg_admin/college_stipend.html', rows=rows, total=total,
                           states=states, kind=kind, state=state, sort=sort, q=q,
                           page=page, pages=pages, per_page=_STIPEND_PER_PAGE,
                           all_names=all_names, active_section='goocampus_in')


@login_required
def college_stipend_detail(master_id):
    """Per-speciality stipend/bond/penalty for one college, for ONE degree family."""
    u = _require_admin()
    if not u:
        flash('Access denied', 'error'); return redirect(url_for('dashboard'))
    kind = _s(request.args.get('kind'))
    if kind not in ('medical', 'dnb'):
        kind = 'medical'
    deg_clause, deg_param = _stipend_family(kind)
    conn = get_db()
    try:
        m = conn.execute("SELECT * FROM pg_college_master WHERE id = ?", [master_id]).fetchone()
        if not m:
            flash('College not found', 'error')
            return redirect(url_for('pg_college_stipend', kind=kind))
        names = [a['alias_name'] for a in conn.execute(
            "SELECT alias_name FROM pg_college_alias WHERE master_id = ?", [master_id]).fetchall()]
        specialities = []
        if names:
            ph = ','.join(['?'] * len(names))
            specialities = conn.execute(
                f"SELECT c.course AS course, MAX(c.degree) AS degree, "
                f"MAX(c.stipend) AS s1, MAX(c.stipend_yr2) AS s2, MAX(c.stipend_yr3) AS s3, "
                f"MAX(c.bond_years) AS bond, MAX(c.penalty) AS penalty "
                f"FROM pg_cutoffs c WHERE c.institute IN ({ph}) AND " + deg_clause +
                " AND (c.stipend IS NOT NULL OR c.bond_years IS NOT NULL OR c.penalty IS NOT NULL) "
                " GROUP BY c.course ORDER BY c.course", names + [deg_param]).fetchall()
    finally:
        conn.close()
    return render_template('pg_admin/college_stipend_detail.html', m=m,
                           specialities=specialities, back_kind=kind, kind=kind,
                           active_section='goocampus_in')
